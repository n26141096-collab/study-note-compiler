from docx.shared import Mm, Pt, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

A4_WIDTH_MM = 210.0
A4_HEIGHT_MM = 297.0


def configure_page(doc, cfg):
    """Apply physical paper size as well as margins.

    v2.1.0 deliberately writes page_width/page_height so a declared A4 profile
    cannot silently remain Word's default US Letter size.
    """
    sec = doc.sections[0]
    page = cfg['page']
    size = str(page.get('size', 'A4')).upper()
    orientation = str(page.get('orientation', 'portrait')).lower()
    if size != 'A4':
        raise ValueError(f'Unsupported page size for release profile: {size}')

    if orientation == 'portrait':
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width = Mm(A4_WIDTH_MM)
        sec.page_height = Mm(A4_HEIGHT_MM)
    elif orientation == 'landscape':
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width = Mm(A4_HEIGHT_MM)
        sec.page_height = Mm(A4_WIDTH_MM)
    else:
        raise ValueError(f'Unsupported orientation: {orientation}')

    m = page['margin_mm']
    sec.top_margin = Mm(m['top'])
    sec.bottom_margin = Mm(m['bottom'])
    sec.left_margin = Mm(m['left'])
    sec.right_margin = Mm(m['right'])


def copy_page_geometry(src, dst):
    dst.page_width = src.page_width
    dst.page_height = src.page_height
    dst.orientation = src.orientation
    dst.top_margin = src.top_margin
    dst.bottom_margin = src.bottom_margin
    dst.left_margin = src.left_margin
    dst.right_margin = src.right_margin
    dst.header_distance = src.header_distance
    dst.footer_distance = src.footer_distance


def set_section_columns(section, num=1, spacing_mm=5.0, separator_line=False):
    sectPr = section._sectPr
    cols = sectPr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols')
        sectPr.append(cols)
    cols.set(qn('w:num'), str(int(num)))
    cols.set(qn('w:space'), str(int(Mm(spacing_mm).twips)))
    cols.set(qn('w:equalWidth'), '1')
    if separator_line:
        cols.set(qn('w:sep'), '1')
    elif qn('w:sep') in cols.attrib:
        del cols.attrib[qn('w:sep')]


def switch_columns(doc, num, cfg):
    current = doc.sections[-1]
    current_cols = current._sectPr.find(qn('w:cols'))
    current_num = 1
    if current_cols is not None and current_cols.get(qn('w:num')):
        current_num = int(current_cols.get(qn('w:num')))
    if current_num == num:
        return current
    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    # Explicitly preserve A4 and margins across continuous section changes.
    copy_page_geometry(current, sec)
    set_section_columns(
        sec,
        num=num,
        spacing_mm=cfg['columns']['spacing_mm'],
        separator_line=cfg['columns']['separator_line'],
    )
    return sec


def add_header_footer(doc, course, topic, cfg):
    for sec in doc.sections:
        header = sec.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.text = ''
        r = p.add_run(cfg['header']['template'].format(course=course, topic=topic))
        r.font.size = Pt(cfg['header']['size_pt'])

        footer = sec.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.text = ''
        r = fp.add_run('Page ')
        r.font.size = Pt(cfg['footer']['size_pt'])
        begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
        instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
        separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
        text = OxmlElement('w:t'); text.text = '1'
        end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
        r._r.extend([begin, instr, separate, text, end])


def set_cell_margins(cell, margin_twips=45):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for edge in ('top', 'start', 'bottom', 'end'):
        el = tcMar.find(qn(f'w:{edge}'))
        if el is None:
            el = OxmlElement(f'w:{edge}')
            tcMar.append(el)
        el.set(qn('w:w'), str(margin_twips))
        el.set(qn('w:type'), 'dxa')



def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = trPr.find(qn('w:cantSplit'))
    if el is None:
        el = OxmlElement('w:cantSplit')
        trPr.append(el)
    el.set(qn('w:val'), 'true')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = trPr.find(qn('w:tblHeader'))
    if el is None:
        el = OxmlElement('w:tblHeader')
        trPr.append(el)
    el.set(qn('w:val'), 'true')


def set_table_borders(table, outer='4', inner='2', center_divider=False):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = borders.find(qn('w:' + edge))
        if el is None:
            el = OxmlElement('w:' + edge)
            borders.append(el)
        if center_divider:
            val = 'single' if edge == 'insideV' else 'nil'
        else:
            val = 'single'
        el.set(qn('w:val'), val)
        el.set(qn('w:sz'), inner if edge.startswith('inside') else outer)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '888888')


def style_table(table, cfg):
    from docx_styles import set_cell_shading
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:type'), 'pct'); tblW.set(qn('w:w'), '5000')
    set_table_borders(table, str(cfg['tables']['outer_border_eighth_points']), str(cfg['tables']['inner_border_eighth_points']))
    if table.rows and cfg['tables']['repeat_header']:
        set_repeat_table_header(table.rows[0])
    keep_limit = int(cfg['tables'].get('keep_small_table_rows_together', 0) or 0)
    keep_whole = keep_limit and len(table.rows) <= keep_limit
    for ri, row in enumerate(table.rows):
        set_row_cant_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell, cfg['tables']['cell_margin_twips'])
            if keep_whole and ri < len(table.rows) - 1:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
            if ri == 0:
                set_cell_shading(cell, cfg['tables']['header_fill'])
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True


def make_two_panel(doc, cfg):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, '0', '4', center_divider=cfg['panels']['center_divider'])
    ratios = [float(cfg['panels']['left_ratio']), float(cfg['panels']['right_ratio'])]
    ratio_total = sum(ratios)
    if ratio_total <= 0:
        raise ValueError('Two-panel ratios must sum to a positive value')
    ratios = [ratio / ratio_total for ratio in ratios]

    section = doc.sections[-1]
    usable_twips = int(
        section.page_width.twips
        - section.left_margin.twips
        - section.right_margin.twips
    )
    widths = [int(usable_twips * ratios[0]), usable_twips - int(usable_twips * ratios[0])]

    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:type'), 'dxa')
    tblW.set(qn('w:w'), str(usable_twips))
    tblLayout = tblPr.find(qn('w:tblLayout'))
    if tblLayout is None:
        tblLayout = OxmlElement('w:tblLayout')
        tblPr.append(tblLayout)
    tblLayout.set(qn('w:type'), 'fixed')

    grid_cols = table._tbl.tblGrid.findall(qn('w:gridCol'))
    for idx, grid_col in enumerate(grid_cols[:2]):
        grid_col.set(qn('w:w'), str(widths[idx]))
    for idx, cell in enumerate(table.rows[0].cells):
        cell.width = Twips(widths[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(cell, 70)
        tcPr = cell._tc.get_or_add_tcPr()
        tcW = tcPr.find(qn('w:tcW'))
        if tcW is None:
            tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
        tcW.set(qn('w:type'), 'dxa')
        tcW.set(qn('w:w'), str(widths[idx]))
    return table


def add_column_break(paragraph):
    run = paragraph.add_run()
    run.add_break(WD_BREAK.COLUMN)
