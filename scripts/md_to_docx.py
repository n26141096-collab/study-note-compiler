#!/usr/bin/env python3
import argparse
import re
import sys
from pathlib import Path
import yaml
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
from docx_styles import apply_base_styles, set_cell_shading, set_paragraph_shading
from docx_layout import configure_page, set_section_columns, switch_columns, add_header_footer, style_table, make_two_panel, add_column_break


def load_cfg(path=None):
    p = Path(path) if path else ROOT / 'config' / 'word_layout.yaml'
    return yaml.safe_load(p.read_text(encoding='utf-8'))


def clean_inline(text):
    return re.sub(r'[`*_]', '', text).strip()


def add_runs(paragraph, text):
    token_re = re.compile(r'(\*\*(.+?)\*\*|`(.+?)`)')
    pos = 0
    for m in token_re.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        if m.group(2) is not None:
            r = paragraph.add_run(m.group(2)); r.bold = True
        else:
            r = paragraph.add_run(m.group(3))
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def parse_md_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].lstrip().startswith('|'):
        cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
        separator = all(re.fullmatch(r':?-{3,}:?', c.replace(' ', '')) for c in cells)
        if not separator:
            rows.append(cells)
        i += 1
    return rows, i


def add_table(container, rows, cfg):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = container.add_table(rows=len(rows), cols=cols)
    for ri, row in enumerate(rows):
        for ci in range(cols):
            val = row[ci] if ci < len(row) else ''
            p = table.cell(ri, ci).paragraphs[0]
            p.style = 'Compact'
            add_runs(p, val)
    style_table(table, cfg)


def add_block(container, line, cfg):
    s = line.rstrip()
    if not s.strip() or s.strip() == '---':
        return
    if re.fullmatch(r'\s*<!--.*?-->\s*', s):
        return
    if s.startswith('# '):
        p = container.add_paragraph(style='Title')
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_runs(p, s[2:])
        set_paragraph_shading(p, cfg['boxes']['core_axis_fill'])
        return
    if s.startswith('## '):
        p = container.add_paragraph(style='Heading 1')
        add_runs(p, s[3:])
        set_paragraph_shading(p, cfg['boxes']['section_fill'])
        return
    if s.startswith('### '):
        p = container.add_paragraph(style='Heading 2')
        add_runs(p, s[4:])
        return
    if s.startswith('#### '):
        p = container.add_paragraph(style='Heading 3')
        add_runs(p, s[5:])
        return
    if s.startswith('> '):
        p = container.add_paragraph(style='KeySentence')
        set_paragraph_shading(p, cfg['boxes']['core_axis_fill'])
        add_runs(p, s[2:])
        return
    if s.startswith('適用：') or s.startswith('來源：'):
        p = container.add_paragraph(style='Meta')
        add_runs(p, s)
        return
    if re.match(r'^(題目|Q)\s*\d+', s):
        p = container.add_paragraph(style='Question')
        add_runs(p, s)
        return
    if s.startswith('答：') or s.startswith('簡答：'):
        p = container.add_paragraph(style='Answer')
        add_runs(p, s)
        return
    if re.match(r'^[-*] ', s):
        p = container.add_paragraph(style='Normal')
        p.paragraph_format.left_indent = Pt(9)
        p.paragraph_format.first_line_indent = Pt(-5)
        p.add_run('• ')
        add_runs(p, s[2:])
        return
    if re.match(r'^\d+\. ', s):
        p = container.add_paragraph(style='Normal')
        add_runs(p, s)
        return
    p = container.add_paragraph(style='Normal')
    add_runs(p, s)


def build(md_text, out_path, course='Study Notes', topic=None, cfg_path=None):
    cfg = load_cfg(cfg_path)
    doc = Document()
    configure_page(doc, cfg)
    apply_base_styles(doc, cfg)
    set_section_columns(doc.sections[0], 1, cfg['columns']['spacing_mm'], cfg['columns']['separator_line'])

    lines = md_text.splitlines()
    detected_title = topic or next((clean_inline(x[2:]) for x in lines if x.startswith('# ')), 'Topic')

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        mcols = re.fullmatch(r'\s*<!--\s*columns:\s*([12])\s*-->\s*', line)
        if mcols:
            switch_columns(doc, int(mcols.group(1)), cfg)
            i += 1
            continue
        if re.fullmatch(r'\s*<!--\s*column-break\s*-->\s*', line):
            p = doc.add_paragraph(style='Normal')
            add_column_break(p)
            i += 1
            continue
        if re.fullmatch(r'\s*<!--\s*page-break\s*-->\s*', line):
            doc.add_page_break()
            i += 1
            continue
        if re.fullmatch(r'\s*<!--\s*layout:\s*two-panel\s*-->\s*', line):
            left, right, target = [], [], None
            i += 1
            while i < len(lines):
                s = lines[i].rstrip()
                if s.strip() == ':::left':
                    target = left
                elif s.strip() == ':::right':
                    target = right
                elif s.strip() == ':::end':
                    break
                elif target is not None:
                    target.append(s)
                i += 1
            panel = make_two_panel(doc, cfg)
            for arr, cell in ((left, panel.cell(0, 0)), (right, panel.cell(0, 1))):
                cell.text = ''
                j = 0
                while j < len(arr):
                    if arr[j].lstrip().startswith('|'):
                        rows, j2 = parse_md_table(arr, j)
                        add_table(cell, rows, cfg)
                        j = j2
                        continue
                    add_block(cell, arr[j], cfg)
                    j += 1
            i += 1
            continue
        if line.lstrip().startswith('|'):
            rows, j = parse_md_table(lines, i)
            add_table(doc, rows, cfg)
            i = j
            continue
        add_block(doc, line, cfg)
        i += 1

    for sec in doc.sections:
        # Same header/footer for every section created by column switches.
        pass
    add_header_footer(doc, course, detected_title, cfg)

    # Paragraph stability, including paragraphs inside tables.
    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True
        if p.style.name in ('Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Question'):
            p.paragraph_format.keep_with_next = True
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.widow_control = True
                    if p.style.name.startswith('Heading') or p.style.name == 'Question':
                        p.paragraph_format.keep_with_next = True

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def main():
    ap = argparse.ArgumentParser(description='Render study-note Markdown to deterministic DOCX.')
    ap.add_argument('input')
    ap.add_argument('--out', required=True)
    ap.add_argument('--course', default='Study Notes')
    ap.add_argument('--topic')
    ap.add_argument('--config')
    args = ap.parse_args()
    text = Path(args.input).read_text(encoding='utf-8')
    print(build(text, args.out, args.course, args.topic, args.config))


if __name__ == '__main__':
    main()
